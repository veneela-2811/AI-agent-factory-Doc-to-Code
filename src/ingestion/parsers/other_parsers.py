import re
from pathlib import Path
from pptx import Presentation
from openpyxl import load_workbook
from src.ingestion.parsers.base import BaseDocumentParser, ParsedDocument, ParsedSection


class PPTXParser(BaseDocumentParser):
    def parse(self, file_path: Path) -> ParsedDocument:
        prs = Presentation(str(file_path))
        sections = []
        raw_parts = []

        for idx, slide in enumerate(prs.slides):
            slide_texts = []
            title = f"Slide {idx + 1}"
            
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        slide_texts.append(text)
            
            if slide_texts:
                title = slide_texts[0][:80]
                content = "\n".join(slide_texts)
                raw_parts.append(content)
                sections.append(ParsedSection(
                    section_id=f"slide_{idx + 1}",
                    title=title,
                    level=1,
                    page_number=idx + 1,
                    content=content
                ))

        return ParsedDocument(
            filename=file_path.name,
            doc_type="PPTX",
            sections=sections,
            raw_text="\n\n".join(raw_parts)
        )


class XLSXParser(BaseDocumentParser):
    def parse(self, file_path: Path) -> ParsedDocument:
        wb = load_workbook(filename=str(file_path), data_only=True)
        sections = []
        raw_parts = []

        for idx, sheetname in enumerate(wb.sheetnames):
            ws = wb[sheetname]
            sheet_rows = []
            for row in ws.iter_rows(values_only=True):
                row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if row_vals:
                    sheet_rows.append(" | ".join(row_vals))

            if sheet_rows:
                content = f"Sheet: {sheetname}\n" + "\n".join(sheet_rows)
                raw_parts.append(content)
                sections.append(ParsedSection(
                    section_id=f"sheet_{idx + 1}",
                    title=f"Sheet - {sheetname}",
                    level=1,
                    content=content
                ))

        return ParsedDocument(
            filename=file_path.name,
            doc_type="XLSX",
            sections=sections,
            raw_text="\n\n".join(raw_parts)
        )


class TextParser(BaseDocumentParser):
    def parse(self, file_path: Path) -> ParsedDocument:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()

        sections = []
        # Parse markdown headings (# , ## , ### )
        heading_regex = re.compile(r"^(#{1,6})\s+(.*)$", re.MULTILINE)
        matches = list(heading_regex.finditer(text))

        if matches:
            for idx, match in enumerate(matches):
                level = len(match.group(1))
                title = match.group(2).strip()
                start_pos = match.end()
                end_pos = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
                sec_content = text[start_pos:end_pos].strip()
                
                sections.append(ParsedSection(
                    section_id=f"sec_{idx + 1}",
                    title=title,
                    level=level,
                    content=sec_content or title
                ))
        else:
            sections.append(ParsedSection(
                section_id="sec_1",
                title="Document Content",
                level=1,
                content=text
            ))

        return ParsedDocument(
            filename=file_path.name,
            doc_type="TEXT",
            sections=sections,
            raw_text=text
        )
