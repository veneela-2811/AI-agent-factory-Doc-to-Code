import pytest
from pathlib import Path
from src.ingestion.parsers import (
    PDFParser, DocxParser, PPTXParser, XLSXParser, TextParser
)
from openpyxl import Workbook
from pptx import Presentation


def test_markdown_text_parser(tmp_path: Path):
    doc_path = tmp_path / "sample.md"
    content = """# System Overview
This is an overview of the agent factory.

## Functional Requirements
- Requirement 1: Ingest documents.
- Requirement 2: Run planning.

### Sub-section
Additional details here.
"""
    doc_path.write_text(content, encoding="utf-8")

    parser = TextParser()
    parsed = parser.parse(doc_path)
    assert parsed.doc_type == "TEXT"
    assert len(parsed.sections) == 3
    assert parsed.sections[0].title == "System Overview"
    assert parsed.sections[1].title == "Functional Requirements"
    assert parsed.sections[2].title == "Sub-section"


def test_xlsx_parser(tmp_path: Path):
    xlsx_path = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"
    ws.append(["ReqID", "Title", "Priority"])
    ws.append(["REQ-1", "Parser support", "High"])
    ws.append(["REQ-2", "Vector indexing", "Medium"])
    wb.save(str(xlsx_path))

    parser = XLSXParser()
    parsed = parser.parse(xlsx_path)
    assert parsed.doc_type == "XLSX"
    assert len(parsed.sections) == 1
    assert "Parser support" in parsed.sections[0].content


def test_pptx_parser(tmp_path: Path):
    pptx_path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Architecture Deck"
    subtitle.text = "Multi-Agent System Overview"
    prs.save(str(pptx_path))

    parser = PPTXParser()
    parsed = parser.parse(pptx_path)
    assert parsed.doc_type == "PPTX"
    assert len(parsed.sections) == 1
    assert "Architecture Deck" in parsed.sections[0].content
