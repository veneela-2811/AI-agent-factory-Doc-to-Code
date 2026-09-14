import pytest
from pathlib import Path
from docx import Document as DocxDoc
from pypdf import PdfWriter
from io import BytesIO
from src.ingestion.parsers import DocxParser, PDFParser


def test_docx_parser(tmp_path: Path):
    docx_file = tmp_path / "sample.docx"
    doc = DocxDoc()
    doc.add_heading("Product Requirements", level=1)
    doc.add_paragraph("This document outlines the core agent capabilities.")
    doc.add_heading("Non-Functional Specs", level=2)
    doc.add_paragraph("System must achieve <500ms latency.")
    doc.save(str(docx_file))

    parser = DocxParser()
    parsed = parser.parse(docx_file)
    assert parsed.doc_type == "DOCX"
    assert len(parsed.sections) == 2
    assert parsed.sections[0].title == "Product Requirements"
    assert parsed.sections[1].title == "Non-Functional Specs"


def test_pdf_parser(tmp_path: Path):
    pdf_file = tmp_path / "sample.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(pdf_file, "wb") as f:
        writer.write(f)

    parser = PDFParser()
    parsed = parser.parse(pdf_file)
    assert parsed.doc_type == "PDF"
